from __future__ import annotations

import argparse
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor import ExtractionResult, extract_pdf


# Keep the original template spelling and column order.
HEADERS = [
    "Plant",
    "Purchae Order",
    "PO/LI",
    "Material",
    "Grid Value",
    "Roll Qty",
    "Vendor Batch number",
    "Vendor Roll Number",
    "Vendor Invoice number",
    "",
    "item",
    "color",
    "N.W",
    "G.W",
]

COLUMN_WIDTHS = [
    12, 20, 12, 18, 16, 14, 24,
    24, 27, 15, 23, 58, 14, 14,
]


def export_excel(
    result: ExtractionResult,
    output_path: str | Path,
) -> Path:
    if not result.rows:
        raise ValueError("No extracted rows are available to export.")

    output_path = Path(output_path).resolve()

    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("The output filename must end with .xlsx")

    if not output_path.parent.is_dir():
        raise ValueError(
            f"Output folder does not exist: {output_path.parent}"
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sap"

    sheet.append(HEADERS)

    for roll in result.rows:
        sheet.append([
            None,                       # A: Plant
            roll.purchase_order,        # B: Purchae Order
            None,                       # C: PO/LI
            None,                       # D: Material
            None,                       # E: Grid Value
            float(roll.quantity),       # F: Roll Qty
            roll.lot,                   # G: Vendor Batch number
            roll.package,               # H: Vendor Roll Number
            roll.invoice,               # I: Vendor Invoice number
            roll.foc_note or None,      # J: Blank heading / FOC note
            roll.item,                  # K: item
            roll.color,                 # L: color
            (
                float(roll.net_weight)
                if roll.net_weight is not None
                else None
            ),
            (
                float(roll.gross_weight)
                if roll.gross_weight is not None
                else None
            ),
        ])

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1D4ED8",
    )
    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF",
    )
    body_font = Font(
        name="Calibri",
        size=11,
        color="172033",
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    sheet.row_dimensions[1].height = 34

    # Includes J even though its heading is intentionally blank.
    sheet.auto_filter.ref = f"A1:N{sheet.max_row}"
    sheet.freeze_panes = "A2"

    for column_number, width in enumerate(COLUMN_WIDTHS, start=1):
        letter = get_column_letter(column_number)
        sheet.column_dimensions[letter].width = width

    text_columns = (2, 7, 8, 9, 10, 11, 12)

    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 30

        for cell in sheet[row_number]:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")

        # Keep identifiers as text, including LOT values such as 2-1.
        # Force strings to remain literal rather than Excel formulas.
        for column_number in text_columns:
            cell = sheet.cell(row_number, column_number)

            if cell.value is not None:
                cell.value = str(cell.value)
                cell.data_type = "s"
                cell.number_format = "@"

        sheet.cell(row_number, 6).number_format = "#,##0.###"
        sheet.cell(row_number, 13).number_format = "#,##0.00"
        sheet.cell(row_number, 14).number_format = "#,##0.00"

        sheet.cell(row_number, 12).alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    sheet.sheet_view.zoomScale = 90
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    # Write and verify a temporary file before replacing the destination.
    descriptor, temporary_name = tempfile.mkstemp(
        suffix=".xlsx",
        prefix="stantex_",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)

    try:
        workbook.save(temporary_path)
        workbook.close()

        verify_export(temporary_path, result)

        os.replace(temporary_path, output_path)

    except PermissionError as error:
        raise PermissionError(
            "Cannot save the Excel file. Close it in Excel if it is "
            "already open, or select a writable output folder."
        ) from error

    finally:
        workbook.close()

        if temporary_path.exists():
            temporary_path.unlink()

    return output_path


def verify_export(
    path: Path,
    result: ExtractionResult,
) -> None:
    workbook = load_workbook(path)

    try:
        sheet = workbook["sap"]

        actual_headers = [
            cell.value if cell.value is not None else ""
            for cell in sheet[1]
        ]

        if actual_headers != HEADERS:
            raise ValueError("Export validation failed: wrong headings.")

        if sheet.max_row != len(result.rows) + 1:
            raise ValueError("Export validation failed: wrong row count.")

        if sheet.auto_filter.ref != f"A1:N{sheet.max_row}":
            raise ValueError("Export validation failed: missing filters.")

        if sheet.freeze_panes != "A2":
            raise ValueError("Export validation failed: header not frozen.")

        for row_number, roll in enumerate(result.rows, start=2):
            expected = [
                None,
                roll.purchase_order,
                None,
                None,
                None,
                float(roll.quantity),
                roll.lot,
                roll.package,
                roll.invoice,
                roll.foc_note or None,
                roll.item,
                roll.color,
                (
                    float(roll.net_weight)
                    if roll.net_weight is not None
                    else None
                ),
                (
                    float(roll.gross_weight)
                    if roll.gross_weight is not None
                    else None
                ),
            ]

            actual = [
                sheet.cell(row_number, column).value
                for column in range(1, 15)
            ]

            if actual != expected:
                raise ValueError(
                    f"Export validation failed at Excel row {row_number}."
                )

            for column in (2, 7, 8, 9, 11, 12):
                if sheet.cell(row_number, column).data_type != "s":
                    raise ValueError(
                        f"Export validation failed: identifier is not "
                        f"text at row {row_number}, column {column}."
                    )

    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert a STANTEX PDF to the Excel template format."
    )
    parser.add_argument("pdf", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="Destination .xlsx file",
    )
    args = parser.parse_args()

    try:
        result = extract_pdf(args.pdf)

        # The desktop UI will allow users to review warnings.
        # This command-line test stops rather than silently exporting them.
        if result.warnings:
            print("\nReview required:")
            for warning in result.warnings:
                print(f"- {warning}")
            print("\nExcel was not exported.")
            return 1

        saved_path = export_excel(result, args.output)

    except Exception as error:
        print(f"\nERROR: {error}")
        return 1

    print("\nEXCEL EXPORT SUCCESS")
    print(f"Rows     : {len(result.rows)}")
    print(f"Quantity : {result.total_quantity:,.0f} YDS")
    print(f"File     : {saved_path}")
    print("Validation: all exported rows match extracted data.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())