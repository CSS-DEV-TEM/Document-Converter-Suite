from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADERS = [
    "Plant",
    "Purchae Order",
    "PO/LI",
    "Material",
    "Grid Value",
    "Roll Qty",
    "Vendor Batch number",
    "Vendor Roll Number",
    "Vendor Invoice number",
    "item",
]


def export_excel(rows, output_path):
    if not rows:
        raise ValueError("No validated rows available to export.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sap"
    sheet.append(HEADERS)

    for row in rows:
        sheet.append([
            None,
            row.purchase_order,
            None,
            None,
            None,
            float(row.quantity),
            row.batch,
            row.vendor_roll,
            row.invoice,
            row.item,
        ])

        # Explicit text cells preserve identifiers and leading zeros.
        # They also prevent identifiers beginning with "=" becoming formulas.
        for column in (2, 7, 8, 9, 10):
            cell = sheet.cell(sheet.max_row, column)
            cell.data_type = "s"
            cell.number_format = "@"

        quantity_cell = sheet.cell(sheet.max_row, 6)
        decimals = max(
            0,
            -row.quantity.normalize().as_tuple().exponent,
        )
        quantity_cell.number_format = (
            "0" if decimals == 0 else "0." + ("0" * decimals)
        )

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    widths = [13, 22, 12, 18, 16, 16, 25, 24, 26, 22]

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"

    workbook.save(output_path)